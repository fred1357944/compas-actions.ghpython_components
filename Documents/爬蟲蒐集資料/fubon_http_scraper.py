#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
Fubon HTTP Scraper - 優化版本
使用直接HTTP請求，支援多時間範圍，高效率爬取富邦證券排行榜資料
"""

import asyncio
import datetime as dt
import json
import re
import time
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import aiofiles
import httpx
import pandas as pd
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from openpyxl.styles import numbers

# ====== 基本配置 ======
BASE_DIR = Path("/Users/laihongyi/Documents/爬蟲蒐集資料")
OUT_DIR = BASE_DIR
SNAP_DIR = BASE_DIR / "_snapshots"
LOG_PATH = BASE_DIR / "_last_run.log"

# 富邦證券排行榜URL配置
FUBON_URLS = {
    "漲幅排行": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZG_A.djhtm",
    "跌幅排行": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZG_B.djhtm", 
    "成交量排行": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZG_D.djhtm",
    "成交值排行": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZG_E.djhtm",
    "當日強勢股": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZGK_DD.djhtm",
    "三日強勢股": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZGK_CC.djhtm",
    "五日強勢股": "https://fubon-ebrokerdj.fbs.com.tw/Z/ZG/ZGK_AA.djhtm",
}

# HTTP請求設定
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/123.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-TW,zh;q=0.9,en;q=0.8",
    "Accept-Encoding": "gzip, deflate, br",
    "Cache-Control": "no-cache",
    "Pragma": "no-cache",
}

# ====== 工具函數 ======
def log(msg: str):
    """記錄日誌"""
    timestamp = dt.datetime.now().strftime('%H:%M:%S')
    msg_with_time = f"[{timestamp}] {msg}"
    print(msg_with_time)
    with open(LOG_PATH, "a", encoding="utf-8") as f:
        f.write(msg_with_time + "\n")

def ensure_dirs():
    """確保目錄存在"""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    SNAP_DIR.mkdir(parents=True, exist_ok=True)

def today_tag() -> str:
    """取得今天日期標籤"""
    return dt.datetime.now().strftime("%Y%m%d")

def extract_stock_data(html_content: str) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    """從HTML內容提取股票資料"""
    soup = BeautifulSoup(html_content, 'lxml')
    
    # 尋找表格
    tables = soup.find_all('table')
    best_table = None
    max_rows = 0
    
    for table in tables:
        rows = table.find_all('tr')
        if len(rows) > max_rows:
            max_rows = len(rows)
            best_table = table
    
    if not best_table:
        # 嘗試從文字中解析股票資料
        return extract_from_text(html_content)
    
    # 解析表格
    data = []
    headers = []
    
    for row in best_table.find_all('tr'):
        cells = row.find_all(['td', 'th'])
        if not cells:
            continue
            
        cell_texts = [cell.get_text(strip=True) for cell in cells]
        
        # 檢查是否為標題行
        if not headers and any('名次' in cell or '代號' in cell or '股票' in cell for cell in cell_texts):
            headers = cell_texts
        else:
            # 過濾空行和無效資料
            if any(cell_texts) and len(cell_texts) >= 4:
                data.append(cell_texts)
    
    if not data:
        return None, None
        
    df = pd.DataFrame(data)
    
    # 設定欄位名稱
    if headers and len(headers) == len(df.columns):
        df.columns = headers
    else:
        # 自動推斷欄位名稱
        cols = []
        for i in range(len(df.columns)):
            if i == 0:
                cols.append('名次')
            elif i == 1:
                cols.append('代號')
            elif i == 2:
                cols.append('股票名稱')
            elif i == 3:
                cols.append('收盤價')
            elif i == 4:
                cols.append('漲跌')
            elif i == 5:
                cols.append('漲跌幅')
            elif i == 6:
                cols.append('成交量')
            else:
                cols.append(f'欄位{i+1}')
        df.columns = cols[:len(df.columns)]
    
    # 提取日期
    date_match = re.search(r'(\d{4}/\d{2}/\d{2})', html_content)
    date_str = date_match.group(1) if date_match else dt.datetime.now().strftime("%Y/%m/%d")
    
    return df, date_str

def extract_from_text(html_content: str) -> Tuple[Optional[pd.DataFrame], Optional[str]]:
    """從純文字內容解析股票資料"""
    soup = BeautifulSoup(html_content, 'lxml')
    text = soup.get_text()
    
    # 尋找股票代碼模式
    stock_pattern = r'(\d+)\s+([0-9A-Z]{4,6})\s*([^\d\n]+?)\s+([0-9,.]+)\s*([+-]?[0-9,.]+)\s*([+-]?[0-9.,%]+)'
    matches = re.findall(stock_pattern, text)
    
    if not matches:
        return None, None
    
    data = []
    for match in matches:
        rank, code, name, price, change, change_pct = match
        data.append([rank.strip(), code.strip(), name.strip(), price.strip(), change.strip(), change_pct.strip()])
    
    if data:
        df = pd.DataFrame(data, columns=['名次', '代號', '股票名稱', '收盤價', '漲跌', '漲跌幅'])
        date_str = dt.datetime.now().strftime("%Y/%m/%d")
        return df, date_str
    
    return None, None

def clean_numeric_value(value) -> Optional[float]:
    """清理數值欄位"""
    if pd.isna(value):
        return None
        
    value_str = str(value).strip()
    if not value_str or value_str in ['--', '-', 'N/A']:
        return None
    
    # 處理萬、億單位
    multiplier = 1
    if '萬' in value_str:
        multiplier = 10000
        value_str = value_str.replace('萬', '')
    elif '億' in value_str:
        multiplier = 100000000
        value_str = value_str.replace('億', '')
    
    # 清理特殊字符
    value_str = (value_str.replace(',', '')
                         .replace('+', '')
                         .replace('＋', '')
                         .replace('－', '-')
                         .replace('%', '')
                         .replace('％', '')
                         .replace('張', '')
                         .replace('元', ''))
    
    try:
        return float(value_str) * multiplier
    except (ValueError, TypeError):
        return None

def process_dataframe(df: pd.DataFrame, source_name: str) -> pd.DataFrame:
    """處理和清理DataFrame"""
    if df is None or df.empty:
        return pd.DataFrame()
    
    df = df.copy()
    
    # 分離代號和股票名稱
    if '代號' not in df.columns and '股票名稱' in df.columns:
        stock_info = df['股票名稱'].astype(str)
        codes = stock_info.str.extract(r'^([0-9A-Z]{4,6})')[0]
        names = stock_info.str.replace(r'^[0-9A-Z]{4,6}\s*', '', regex=True)
        
        df.insert(1, '代號', codes)
        df['股票名稱'] = names
    
    # 清理數值欄位
    numeric_cols = ['名次', '收盤價', '漲跌', '漲跌幅', '成交量', '成交值']
    for col in numeric_cols:
        if col in df.columns:
            df[col] = df[col].apply(clean_numeric_value)
    
    # 過濾有效資料
    if '代號' in df.columns:
        df = df[df['代號'].notna() & (df['代號'].astype(str).str.len() >= 4)]
    
    # 添加來源資訊
    df['資料來源'] = source_name
    df['抓取時間'] = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    return df

async def fetch_url(client: httpx.AsyncClient, url: str, name: str) -> Tuple[str, Optional[pd.DataFrame], Optional[str]]:
    """非同步抓取單個URL"""
    try:
        log(f"抓取 {name}: {url}")
        response = await client.get(url, timeout=30.0)
        response.raise_for_status()
        
        df, date_str = extract_stock_data(response.text)
        if df is not None and not df.empty:
            df = process_dataframe(df, name)
            log(f"✅ {name}: 取得 {len(df)} 筆資料")
            
            # 保存快照
            snapshot_path = SNAP_DIR / f"{today_tag()}_{name}.html"
            await save_snapshot(snapshot_path, response.text)
            
            return name, df, date_str
        else:
            log(f"⚠️ {name}: 未找到有效資料")
            return name, None, None
            
    except Exception as e:
        log(f"❌ {name}: 錯誤 - {str(e)}")
        return name, None, None

async def save_snapshot(path: Path, content: str):
    """非同步保存快照"""
    async with aiofiles.open(path, 'w', encoding='utf-8') as f:
        await f.write(content)

async def main():
    """主要執行函數"""
    ensure_dirs()
    
    try:
        LOG_PATH.unlink()
    except:
        pass
    
    log("開始抓取富邦證券排行榜資料...")
    
    all_data = {}
    
    # 設定HTTP客戶端（跳過SSL驗證）
    limits = httpx.Limits(max_keepalive_connections=5, max_connections=10)
    async with httpx.AsyncClient(headers=HEADERS, limits=limits, verify=False) as client:
        # 建立任務列表
        tasks = []
        for name, url in FUBON_URLS.items():
            tasks.append(fetch_url(client, url, name))
        
        # 並行執行
        results = await asyncio.gather(*tasks, return_exceptions=True)
        
        # 處理結果
        for result in results:
            if isinstance(result, Exception):
                log(f"任務執行錯誤: {result}")
                continue
                
            name, df, date_str = result
            if df is not None and not df.empty:
                all_data[name] = df
    
    if not all_data:
        log("❌ 未取得任何資料")
        return
    
    # 生成輸出檔案
    tag = today_tag()
    output_path = OUT_DIR / f"fubon_data_{tag}_enhanced.xlsx"
    
    # 寫入Excel
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        summary_data = []
        
        for sheet_name, df in all_data.items():
            # 限制工作表名稱長度
            safe_sheet_name = sheet_name[:31]
            df.to_excel(writer, sheet_name=safe_sheet_name, index=False)
            
            # 收集彙總資料
            if not df.empty:
                summary_df = df.copy()
                summary_df['分類'] = sheet_name
                summary_data.append(summary_df)
        
        # 建立彙總表
        if summary_data:
            summary = pd.concat(summary_data, ignore_index=True)
            summary.to_excel(writer, sheet_name='彙總資料', index=False)
    
    # 格式化Excel
    format_excel(output_path)
    
    log(f"✅ 完成！輸出檔案: {output_path}")
    log(f"總計取得 {sum(len(df) for df in all_data.values())} 筆資料")

def format_excel(file_path: Path):
    """格式化Excel檔案"""
    try:
        wb = load_workbook(file_path)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            
            # 數字格式化
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if cell.value is not None:
                        # 價格和金額欄位
                        if any(keyword in str(ws.cell(1, cell.column).value or '') 
                               for keyword in ['價', '金額', '成交值']):
                            cell.number_format = '#,##0.00'
                        # 百分比欄位
                        elif '幅' in str(ws.cell(1, cell.column).value or ''):
                            cell.number_format = '0.00%'
                        # 數量欄位
                        elif any(keyword in str(ws.cell(1, cell.column).value or '')
                                for keyword in ['量', '張']):
                            cell.number_format = '#,##0'
        
        wb.save(file_path)
    except Exception as e:
        log(f"Excel格式化錯誤: {e}")

if __name__ == "__main__":
    try:
        asyncio.run(main())
    except KeyboardInterrupt:
        log("程式被使用者中斷")
    except Exception as e:
        log(f"執行錯誤: {e}")
        import traceback
        traceback.print_exc()